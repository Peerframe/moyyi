#!/usr/bin/env python3
"""Generate public, complete document packages and reopen them with independent readers."""

import argparse
import datetime
import getpass
import hashlib
import importlib.metadata
import json
from pathlib import Path, PurePosixPath
import platform
import posixpath
import shutil
import subprocess
import tempfile
from urllib.parse import unquote
import zipfile
import zlib

from docx import Document
from lxml import etree
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pypdf import PdfReader
from reportlab.pdfgen import canvas

HERE = Path(__file__).resolve().parent
FIXED = datetime.datetime(2000, 1, 1)
STAMP = "2000-01-01T00:00:00Z"
VERSIONS = {
    "python-docx": "1.2.0",
    "openpyxl": "3.1.5",
    "python-pptx": "1.0.2",
    "reportlab": "4.4.9",
    "pypdf": "6.10.0",
    "lxml": "6.1.1",
}
LIBREOFFICE = "LibreOfficeDev 26.8.0.0.alpha0 2c87e51eeaa2b413ff4ae097b2705eea1995d8e5"
ODF = {"odt": "docx", "ods": "xlsx", "odp": "pptx"}
EXPORT = {"odt": "docx:Office Open XML Text", "ods": "xlsx:Calc MS Excel 2007 XML", "odp": "pptx:Impress MS PowerPoint 2007 XML"}


def properties(obj):
    obj.title = "moyyi public document fixture"
    obj.subject = "Synthetic parser acceptance input"
    obj.author = "moyyi"
    obj.last_modified_by = "moyyi"
    obj.created = FIXED
    obj.modified = FIXED
    obj.revision = 1


def create(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    extension = path.suffix[1:]
    if extension == "pdf":
        pdf = canvas.Canvas(str(path), pagesize=(612, 792), invariant=1, pageCompression=1)
        pdf.setTitle("moyyi public document fixture")
        pdf.setAuthor("moyyi")
        pdf.setCreator("moyyi fixture generator; ReportLab 4.4.9")
        pdf.setFont("Helvetica", 10)
        pdf.drawString(40, 740, text)
        pdf.showPage()
        pdf.save()
    elif extension == "docx":
        doc = Document()
        properties(doc.core_properties)
        doc.add_paragraph(text)
        doc.save(path)
    elif extension == "xlsx":
        book = Workbook()
        book.properties.creator = "moyyi"
        book.properties.lastModifiedBy = "moyyi"
        book.properties.title = "moyyi public document fixture"
        book.properties.created = FIXED
        book.properties.modified = FIXED
        book.active.title = "Evidence"
        book.active["A1"] = text
        book.active.column_dimensions["A"].width = 110
        book.save(path)
    elif extension == "pptx":
        deck = Presentation()
        properties(deck.core_properties)
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(0.4), Inches(1), Inches(9), Inches(1.5)).text_frame
        frame.text = text
        frame.paragraphs[0].font.size = Pt(16)
        deck.save(path)
    else:
        raise ValueError(extension)


def normalize(path):
    if path.suffix == ".pdf":
        return
    with zipfile.ZipFile(path) as archive:
        entries = {info.filename: archive.read(info) for info in archive.infolist()}
    if "docProps/core.xml" in entries:
        tree = etree.fromstring(entries["docProps/core.xml"])
        for node in tree:
            if etree.QName(node).localname in ("created", "modified", "lastPrinted"):
                node.text = STAMP
        entries["docProps/core.xml"] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
    if "meta.xml" in entries:
        tree = etree.fromstring(entries["meta.xml"])
        for node in tree.iter():
            if etree.QName(node).localname in ("creation-date", "date", "print-date"):
                node.text = STAMP
            elif etree.QName(node).localname in ("initial-creator", "creator", "printed-by"):
                node.text = "moyyi"
            elif etree.QName(node).localname == "editing-duration":
                node.text = "PT0S"
        entries["meta.xml"] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8")
    if "settings.xml" in entries:
        tree = etree.fromstring(entries["settings.xml"])
        for node in tree.iter():
            if node.get("{urn:oasis:names:tc:opendocument:xmlns:config:1.0}name") in ("Rsid", "RsidRoot"):
                node.text = "1"
        entries["settings.xml"] = etree.tostring(tree, xml_declaration=True, encoding="UTF-8")
    # Fixed timestamps/order/permissions remove archive and clock variation. ODF requires
    # uncompressed mimetype first; package contents and relationships otherwise stay intact.
    with zipfile.ZipFile(path, "w") as archive:
        for name in sorted(entries, key=lambda value: (value != "mimetype", value)):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.create_system = 0
            info.external_attr = 0x20
            archive.writestr(info, entries[name], compress_type=zipfile.ZIP_STORED if name == "mimetype" else zipfile.ZIP_DEFLATED, compresslevel=9)


def convert(soffice, source, target, extension, temporary):
    target.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=temporary, prefix="office-profile-") as profile:
        process = subprocess.run(
            [soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--norestore", "--nodefault", "--nofirststartwizard", "--convert-to", extension, "--outdir", str(target), str(source)],
            check=True, text=True, capture_output=True, timeout=45,
            env={"HOME": profile, "TMPDIR": str(temporary), "PATH": "/usr/bin:/bin", "LANG": "C", "LC_ALL": "C"},
        )
        output = target / f"{source.stem}.{extension.split(':')[0]}"
        if not output.is_file():
            raise RuntimeError(f"LibreOffice did not convert {source.name}: {process.stdout} {process.stderr}")
        return output


def read_text(path):
    extension = path.suffix[1:]
    if extension == "pdf":
        pdf = PdfReader(path, strict=True)
        assert len(pdf.pages) == 1
        return "\n".join(page.extract_text() for page in pdf.pages)
    if extension == "docx":
        return "\n".join(p.text for p in Document(path).paragraphs)
    if extension == "xlsx":
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            return "\n".join(str(cell) for sheet in book for row in sheet.values for cell in row if cell is not None)
        finally:
            book.close()
    if extension == "pptx":
        return "\n".join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)
    raise ValueError(extension)


def check_package(path):
    if path.suffix == ".pdf":
        return
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
        names = set(archive.namelist())
        for name in names:
            content = archive.read(name)
            # Fixtures may be committed publicly. Reject machine/user paths even in ancillary parts.
            for private in (b"/Users/", getpass.getuser().encode(), b"file:///", b"/private/var/", b"/tmp/"):
                assert private not in content, f"Private path in {path.name}:{name}"
            if name.endswith((".xml", ".rels")):
                etree.fromstring(content, etree.XMLParser(resolve_entities=False, no_network=True))
            if name.endswith(".rels"):
                base = "" if name == "_rels/.rels" else name.rsplit("/_rels/", 1)[0]
                for rel in etree.fromstring(content):
                    if rel.get("TargetMode") == "External":
                        raise AssertionError(f"External relationship in {path.name}")
                    target = unquote(rel.get("Target").split("#")[0])
                    resolved = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(base, target))
                    assert resolved in names, f"Missing relationship target {resolved} in {path.name}"
        if path.suffix[1:] in ODF:
            first = archive.infolist()[0]
            assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
            assert {"META-INF/manifest.xml", "content.xml", "styles.xml", "meta.xml"} <= names
            manifest = etree.fromstring(archive.read("META-INF/manifest.xml"))
            for entry in manifest:
                name = entry.get("{urn:oasis:names:tc:opendocument:xmlns:manifest:1.0}full-path")
                assert name == "/" or name in names or any(part.startswith(name) for part in names)
        else:
            assert {"[Content_Types].xml", "_rels/.rels", "docProps/core.xml"} <= names
            if path.suffix == ".pptx":
                tree = etree.fromstring(archive.read("ppt/presentation.xml"))
                assert tree.find("{*}sldIdLst/{*}sldId") is not None


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--soffice", default=shutil.which("soffice"))
    parser.add_argument("--output", type=Path, default=HERE)
    parser.add_argument("--verify", action="store_true", help="Reopen and hash existing fixtures without replacing them")
    args = parser.parse_args()
    assert args.soffice, "LibreOffice soffice executable is required"
    versions = {name: importlib.metadata.version(name) for name in VERSIONS}
    assert versions == VERSIONS, f"Use the pinned generator/reader versions: {versions}"
    with tempfile.TemporaryDirectory(prefix="moyyi-office-version-") as profile:
        version = subprocess.check_output(
            [args.soffice, f"-env:UserInstallation={Path(profile).as_uri()}", "--version"],
            text=True, timeout=30,
            env={"HOME": profile, "TMPDIR": profile, "PATH": "/usr/bin:/bin", "LANG": "C", "LC_ALL": "C"},
        ).strip()
    assert version == LIBREOFFICE, f"Use the recorded LibreOffice build: {version}"
    inputs = json.loads((HERE / "inputs.json").read_text())
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    manifest = {"schema": 1, "authorship": "Independently authored public synthetic moyyi fixtures; no private or downloaded documents.", "python": platform.python_version(), "packages": versions, "libreoffice": version, "zlib": zlib.ZLIB_RUNTIME_VERSION, "fixedMetadata": STAMP, "files": {}}
    with tempfile.TemporaryDirectory(prefix="moyyi-standard-documents-") as temp:
        temporary = Path(temp)
        for relative, text in inputs.items():
            path = output / relative
            extension = path.suffix[1:]
            if not args.verify:
                if extension in ODF:
                    seed = temporary / "seed" / f"{PurePosixPath(relative).as_posix().replace('/', '-')}.{ODF[extension]}"
                    create(seed, text)
                    normalized = convert(args.soffice, seed, temporary / "odf", extension, temporary)
                    path.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(normalized, path)
                else:
                    create(path, text)
                normalize(path)
            check_package(path)
            checked = convert(args.soffice, path, temporary / "reopened" / relative.replace("/", "-"), EXPORT[extension], temporary) if extension in ODF else path
            assert text in read_text(checked), f"Mature reader lost expected text in {relative}"
            data = path.read_bytes()
            manifest["files"][relative] = {"sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data), "text": text, "reader": f"LibreOffice round-trip to {ODF[extension]}, then Python reader" if extension in ODF else {"pdf": "pypdf strict=True", "docx": "python-docx", "xlsx": "openpyxl", "pptx": "python-pptx"}[extension]}
            print(f"verified {relative}", flush=True)
    manifest_file = output / "manifest.json"
    if args.verify:
        assert manifest["files"] == json.loads(manifest_file.read_text())["files"], "Fixture hash/reader evidence changed"
    else:
        manifest_file.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n")


if __name__ == "__main__":
    main()
